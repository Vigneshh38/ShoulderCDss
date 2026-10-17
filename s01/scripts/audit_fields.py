"""Audit: is every routing question a real Medical Record Library column, and every option a real value?

    python scripts/audit_fields.py [report.txt]

Reads the questions from src/ShoulderRouter/routingSpec.js (through node) and every row of the 18 Golden Datasets,
then checks, for every question:
  - each [wrapper, column] exists in that wrapper's MRL header;
  - every value that occurs in the data is one of the app's options (nothing in the data the app cannot show);
  - every option occurs in the data, or else is a value the Layer 1 validator allows;
and that every surgery procedure text occurs in the "Previous Shoulder Surgery" cells exactly as spelled.
Writes the full report (default: field_audit.txt) and exits 1 when any check fails.
"""
import collections
import glob
import json
import os
import subprocess
import sys

import openpyxl

S01 = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
ROOT = os.path.dirname(S01)
DUMP = r"""
import { QUESTIONS, CORE_FIELDS, HISTORY_FIELDS, PROCEDURES } from './src/ShoulderRouter/routingSpec.js';
const ov = o => (typeof o === 'object' ? o.v : o);
console.log(JSON.stringify({
  qs: Object.values(QUESTIONS).map(q => ({ key: q.key, ui: !!q.ui, kind: q.kind, cols: q.cols || [], options: (q.options || []).map(ov), describe: !!q.describe })),
  shared: [...CORE_FIELDS, ...HISTORY_FIELDS].map(f => ({ key: f.key, kind: f.kind, options: f.options || [] })),
  procedures: PROCEDURES.map(p => p.v),
}));
"""
SURGERY = 'Previous Shoulder Surgery - Any (Y/N, procedure/date if known)'


def main():
    out_path = sys.argv[1] if len(sys.argv) > 1 else os.path.join(S01, 'field_audit.txt')
    spec = json.loads(subprocess.run(['node', '--input-type=module', '-e', DUMP], cwd=S01, check=True, capture_output=True, text=True, encoding='utf8').stdout)
    lib = openpyxl.load_workbook(os.path.join(ROOT, 'logic_docs', 'Shoulder_CDSS_Field_Rule_Logic_Library.xlsx'), read_only=True)
    l1 = {(r[1], r[3]): str(r[5] or '') for r in list(lib['Layer1_Field_Validators'].iter_rows(values_only=True))[1:] if r[2] == 'MRL'}

    need = collections.defaultdict(set)
    for q in spec['qs']:
        for w, c in q['cols']:
            need[w].add(c)
    shared = [f['key'] for f in spec['shared']]
    headers, values = {}, {}
    for w in [f'S{i:02d}' for i in range(1, 19)]:
        path = glob.glob(os.path.join(ROOT, 'wrappers', w, '*Golden_Dataset*.xlsx'))[0]
        it = openpyxl.load_workbook(path, read_only=True)['Medical Record Library'].iter_rows(values_only=True)
        h = [str(x) for x in next(it)]
        headers[w] = h
        idx = {c: h.index(c) for c in need[w] | set(shared) if c in h}
        cnt = {c: collections.Counter() for c in idx}
        for r in it:
            if r[0] is None:
                continue
            for c, i in idx.items():
                cnt[c][str(r[i]) if r[i] is not None else ''] += 1
        values[w] = cnt
        print(f'read {w}', file=sys.stderr, flush=True)

    lines, problems = [], []
    P = lines.append
    for q in spec['qs']:
        if q['ui']:
            P(f"[{q['key']}] not an Excel column (routing helper)")
            continue
        for w, c in q['cols']:
            if c not in headers[w]:
                problems.append(f'{w} | {c}: not in the MRL header')
                continue
            cnt = values[w][c]
            P(f"[{w} | {c}] {sum(cnt.values())} rows, {len(cnt)} distinct values; Layer 1: {l1.get((w, c), '')[:160]}")
            if not q['options']:
                continue
            ok = lambda v: v in q['options'] or (q['describe'] and v.startswith('Present - '))
            for v, n in cnt.items():
                if not ok(v):
                    problems.append(f'{w} | {c}: data value {v!r} ({n} rows) is not an option')
            for o in q['options']:
                n = cnt.get(o, 0) + (sum(m for v, m in cnt.items() if v.startswith('Present - ')) if q['describe'] and o == 'Present' else 0)
                in_l1 = o.split(' (')[0] in l1.get((w, c), '')
                P(f'    option {o!r}: {n} rows' + ('' if n else ' (none in the data; ' + ('a Layer 1 value)' if in_l1 else 'NOT in Layer 1)')))
                if not n and not in_l1:
                    problems.append(f'{w} | {c}: option {o!r} is neither in the data nor in Layer 1')
    for f in spec['shared']:
        missing = [w for w in headers if f['key'] not in headers[w]]
        if missing:
            problems.append(f"{f['key']}: missing from {missing}")
    surgery = collections.Counter()
    for w in headers:
        surgery.update(values[w].get(SURGERY, {}))
    P('\nProcedures (exact text after "Yes - "):')
    for p in spec['procedures']:
        n = sum(m for v, m in surgery.items() if v.startswith('Yes - ' + p))
        P(f'    {p!r}: {n} rows')
        if not n and p != 'Other (specify)':
            problems.append(f'procedure {p!r} does not occur in any surgery cell')

    P('\nPROBLEMS: ' + (str(len(problems)) if problems else 'none'))
    lines.extend('  ' + x for x in problems)
    with open(out_path, 'w', encoding='utf8') as fh:
        fh.write('\n'.join(lines) + '\n')
    print(f"{len([q for q in spec['qs'] if not q['ui']])} questions, {len(shared)} shared fields, {len(spec['procedures'])} procedures checked; "
          f"{len(problems)} problem(s). Report: {out_path}")
    for x in problems[:20]:
        print('  ' + x)
    sys.exit(1 if problems else 0)


if __name__ == '__main__':
    main()
