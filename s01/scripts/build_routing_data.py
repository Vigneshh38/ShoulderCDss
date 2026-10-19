"""Build src/ShoulderRouter/routingData.generated.js from the project's source files.

Reads:
  wrappers/Sxx/*_Golden_Dataset*.xlsx            Medical Record Library sheet of all 18 wrappers
  wrappers/Sxx/Sxx_Field_Rule_Logic_Document.docx  section 10 (Red Flag Master List) + field notes
  logic_docs/Shoulder_CDSS_Field_Rule_Logic_Library.xlsx   Layer1 validators, DDX rules, grading scales
  logic_docs/Shoulder_CDSS_Workflow_and_Navigation_Guide.docx  branch order + presenting-picture table

Usage (from the s01 folder):
  python scripts/build_routing_data.py
  python scripts/build_routing_data.py --test-out ../tmp/routing_test.json --test-rows 400
  python scripts/build_routing_data.py --test-out ../tmp/routing_all.jsonl --test-rows 0     (every row, one per line)

ROUTING_FIELDS below must list every wrapper-specific column that src/ShoulderRouter/routingSpec.js asks.
"""
import argparse
import collections
import datetime
import glob
import json
import os
import re
import statistics
import sys
import zipfile

import openpyxl

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
WRAPPERS_DIR = os.path.join(ROOT, 'wrappers')
DOCS_DIR = os.path.join(ROOT, 'logic_docs')
OUT = os.path.join(ROOT, 's01', 'src', 'ShoulderRouter', 'routingData.generated.js')

IDS = [f'S{i:02d}' for i in range(1, 19)]

# The 23 shared intake columns (Case ID is system-assigned and not asked). Same header in all 18 workbooks.
SHARED = [
    'Age', 'Gender', 'BMI', 'Co-morbidities', 'Medication History', 'Occupation', 'Activity Level', 'Laterality',
    'Pain_Location', 'Pain_Nature', 'Pain_Intensity', 'Onset',
    'Dominant Arm (Right/Left/Ambidextrous)', 'Affected Side Is Dominant (Y/N)',
    'Previous Shoulder Surgery - Any (Y/N, procedure/date if known)', 'Previous Episode(s) of Same Complaint (Y/N)',
    'Previous Treatment Received (physiotherapy/injection/surgery/none)', 'Aggravating Factors', 'Easing Factors',
    'Night Pain (Y/N)', 'Sleep Disturbance Due to Pain (Y/N)', 'Smoking Status',
    'Cervical Spine Screen (Clear/Not Clear - referred-pain differential)',
]
RED_FLAGS_COL = 'Red Flags / Precautions'
TREATMENT = 'Previous Treatment Received (physiotherapy/injection/surgery/none)'
CERVICAL = 'Cervical Spine Screen (Clear/Not Clear - referred-pain differential)'
FREE_TEXT = ['Medication History', 'Occupation', 'Aggravating Factors', 'Easing Factors']

# Wrapper-specific columns asked by the red-flag screen, the branch walk and the node walk (see routingSpec.js).
ROUTING_FIELDS = {
    'S01': ['Painful Arc (60-120 deg)', 'Empty Can (Jobe) Test', 'Full Can Test', 'Lift-Off Test (subscapularis)',
            "Speed's Test (biceps)", 'Hawkins-Kennedy Impingement Test', 'Neer Impingement Test',
            'External Rotation Lag Sign', 'Drop Arm Test', 'Belly Press Test', "Yergason's Test (biceps)",
            'Muscle Strength (MMT) - ER/IR/Abduction',
            'MRI_Rotator Cuff Tear Grade (partial/full-thickness) / Goutallier Fatty Infiltration Grade',
            'Ultrasound_Tendon Tear Size'],
    'S02': ['Number of Prior Dislocation Episodes', 'Apprehension Test', 'Relocation Test', 'Sulcus Sign',
            "O'Brien's Active Compression Test (SLAP)", 'Jerk Test (posterior instability)',
            'Kim Test (posterior labral)', 'Load and Shift Test', 'Anterior/Posterior Drawer Test (GH joint)',
            'MRI Arthrogram_Labral Tear Location/Type (Bankart / SLAP I-IV / reverse Bankart)',
            'MRI_Hill-Sachs Lesion Present (Y/N)'],
    'S03': ['Painful Arc', 'Hawkins-Kennedy Test', 'Neer Test', 'Subacromial Tenderness (Palpation)',
            'Ultrasound_Subacromial Bursal Thickening/Effusion', 'Response to Diagnostic Subacromial Injection (Y/N)'],
    'S04': ['AC Joint Tenderness/Step Deformity', 'Cross-Body Adduction Test', "O'Brien's Test (AC-specific)",
            'AC Joint Compression Test', 'Piano Key Sign', 'X-Ray_Rockwood Classification (Type I-VI)',
            'Response to Diagnostic AC Joint Injection (Y/N)'],
    'S05': ['Passive ROM Restriction (capsular pattern: ER > Abduction > IR)', 'End-Feel (capsular)'],
    'S06': ['Crepitus', 'X-Ray_Kellgren-Lawrence Grade (Glenohumeral)',
            'Symmetric Joint Involvement (inflammatory screen)',
            'Extra-Articular Manifestations (if inflammatory - nodules/eye/skin/lung)'],
    'S07': ['Neurovascular Status (retrosternal compromise screen - critical for posterior dislocation)',
            'SC Joint Tenderness/Swelling/Deformity', 'Serial Compression Test', 'X-Ray_Serendipity View Findings'],
    'S08': ['Gross Deformity', 'X-Ray_Neer Classification', 'Skin Integrity (open vs closed)',
            'Neurovascular Status (axillary nerve/axillary artery)'],
    'S09': ['Deformity/Shortening', 'X-Ray_Allman Classification', 'Skin Tenting/Open Fracture Screen'],
    'S10': ['Mechanism/Onset (high-energy trauma for fracture / repetitive for scapulothoracic disorders)',
            'Neurovascular Status', 'CT_Scapular Fracture Pattern (Ideberg Classification for glenoid)',
            'Scapulothoracic Crepitus/Snapping', 'Scapular Winging'],
    'S11': ['Scapular Dyskinesis Test Grade', 'Onset (insidious; often overhead athlete)'],
    'S12': ['Procedure Type (Arthroscopic RCR/Open RCR/Bankart Repair/Latarjet/SLAP Repair/Subacromial Decompression)',
            'Post-Operative Week/Phase'],
    'S13': ['Type of Arthroplasty (Anatomic/Reverse/Hemiarthroplasty)', 'Fever/Systemic Illness',
            'Wound Discharge/Erythema', 'Joint Aspirate_Appearance'],
    'S14': ['Sensory Distribution (ulnar nerve/C8-T1)', 'Roos Test (EAST)', "Adson's Test", "Wright's Test (Hyperabduction)",
            'Costoclavicular Maneuver', 'Upper Limb Tension Test (ULTT)', 'MRI/CT_Brachial Plexus Compression Site',
            'Cervical Rib Palpation', 'X-Ray_Cervical Rib Present'],
    'S15': ['Onset (acute arm swelling/discolouration - effort thrombosis - vs chronic intermittent claudication-type)',
            'Radial Pulse Comparison (bilateral)', 'Limb Colour/Temperature Asymmetry',
            'Venous Distension/Swelling (effort thrombosis screen)', 'Limb Girth Comparison (cm, bilateral)'],
    'S16': ['Mechanism/Onset (traumatic traction injury / iatrogenic / compressive)',
            'Axillary Nerve_Deltoid Strength/Sensation (regimental badge area)',
            'Suprascapular Nerve_Infraspinatus/Supraspinatus Strength',
            'Long Thoracic Nerve_Serratus Anterior Strength/Winging', 'Spinal Accessory Nerve_Trapezius Strength',
            'Sensory Mapping (dermatomal)', 'Reflexes (biceps/triceps)', 'EMG/Nerve Conduction Study_Findings',
            'MRI_Nerve Compression Site (suprascapular/spinoglenoid notch)'],
    'S17': ['Onset (birth-related / developmental)', 'Birth History (shoulder dystocia/traction injury)',
            'X-Ray_Os Acromiale Present', 'Os Acromiale Palpation'],
    'S18': ['MRI/Biopsy_Suspected Tumour Findings', 'Crepitus/Mass Palpation', 'Symptom Pattern (atypical/non-mechanical)',
            'Referred Pain Screen (cervical spine/cardiac/diaphragmatic)'],
}

TOP_N = 10


def docx_paragraphs(path):
    xml = zipfile.ZipFile(path).read('word/document.xml').decode('utf8')
    out = []
    for p in re.findall(r'<w:p[ >].*?</w:p>', xml, flags=re.S):
        t = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', p))
        t = t.replace('&amp;', '&').replace('&lt;', '<').replace('&gt;', '>').replace('&quot;', '"').replace('&apos;', "'")
        if t.strip():
            out.append(t.strip())
    return out


def surgery_category(v):
    s = str(v or '')
    if s.startswith('No') or not s:
        return 'No'
    if 'index procedure for this record' in s:
        return 'Yes - index procedure for this record'
    if 'for this fracture' in s:
        return 'Yes - fixation for this fracture'
    return 'Yes - other / earlier surgery'


def top(counter, n=TOP_N):
    return [[k, c] for k, c in counter.most_common(n)]


def read_library():
    wb = openpyxl.load_workbook(os.path.join(DOCS_DIR, 'Shoulder_CDSS_Field_Rule_Logic_Library.xlsx'), read_only=True)
    l1 = {}
    for r in list(wb['Layer1_Field_Validators'].iter_rows(values_only=True))[1:]:
        if r[2] == 'MRL':
            l1[(r[1], r[3])] = {'rule': r[0], 'type': r[4], 'values': r[5], 'notes': r[6]}
    ddx = collections.defaultdict(list)
    for r in list(wb['Layer3_DDX_Rules'].iter_rows(values_only=True))[1:]:
        if r[1]:
            ddx[r[1]].append({'rule': r[0], 'name': r[2], 'age': r[3], 'onset': r[4], 'ddx': r[5]})
    grading = {}
    for r in list(wb['Grading_Staging_Scale_Reference'].iter_rows(values_only=True))[1:]:
        if r[0]:
            grading[r[0]] = {'name': r[1], 'field': r[2], 'values': r[3]}
    return l1, ddx, grading


def read_red_flags():
    """Section 10 of each wrapper's Field Rule Logic Document lists the SH-RF rules that route through it."""
    flags, documented = {}, {}
    for w in IDS:
        doc = glob.glob(os.path.join(WRAPPERS_DIR, w, '*Field_Rule_Logic_Document.docx'))[0]
        ps = docx_paragraphs(doc)
        try:
            i = next(k for k, p in enumerate(ps) if p.startswith('10. Red Flag Master List'))
        except StopIteration:
            continue
        j = i + 1
        while j < len(ps) and not re.match(r'^11\.', ps[j]):
            if re.fullmatch(r'SH-RF-0\d', ps[j]) and j + 2 < len(ps) and not re.fullmatch(r'SH-RF-0\d', ps[j + 1]):
                fid = ps[j]
                f = flags.setdefault(fid, {'id': fid, 'desc': ps[j + 1], 'urgency': ps[j + 2], 'wrappers': []})
                if w not in f['wrappers']:
                    f['wrappers'].append(w)
                j += 3
                continue
            j += 1
        # Field-level notes such as S07: 'A "Not Clear" result here is red flag SH-RF-01 (Section 10)'.
        for k, p in enumerate(ps):
            m = re.search(r'is red flag (SH-RF-0\d)', p)
            if m:
                documented.setdefault(m.group(1), []).append({'wrapper': w, 'text': p})
    return dict(sorted(flags.items())), documented


def read_workflow():
    ps = docx_paragraphs(os.path.join(DOCS_DIR, 'Shoulder_CDSS_Workflow_and_Navigation_Guide.docx'))
    step2 = next(p for p in ps if p.startswith('Step 2'))
    i = next(k for k, p in enumerate(ps) if p.startswith('...consider this group'))
    table = []
    k = i + 1
    while k + 1 < len(ps) and re.search(r'\bS\d\d\b', ps[k + 1]):
        table.append({'picture': ps[k], 'group': ps[k + 1], 'ids': re.findall(r'S\d\d', ps[k + 1])})
        k += 2
    step1 = next(p for p in ps if p.startswith('Step 1'))
    return {'step1': step1, 'step2': step2, 'presenting': table}


def read_dataset(w, test_rows):
    path = glob.glob(os.path.join(WRAPPERS_DIR, w, '*Golden_Dataset*.xlsx'))[0]
    ws = openpyxl.load_workbook(path, read_only=True)['Medical Record Library']
    it = ws.iter_rows(values_only=True)
    header = [str(h) for h in next(it)]
    col = {h: i for i, h in enumerate(header)}
    routing = [c for c in ROUTING_FIELDS.get(w, []) if c in col]
    missing = [c for c in ROUTING_FIELDS.get(w, []) if c not in col]
    if missing:
        sys.exit(f'{w}: routing columns not in the MRL sheet: {missing}')
    counts = {c: collections.Counter() for c in routing + ['Onset', 'Pain_Location', 'Laterality', 'Diagnosis', RED_FLAGS_COL,
                                                          'Co-morbidities', TREATMENT, CERVICAL] + FREE_TEXT}
    surgery = collections.Counter()
    ages = []
    samples, seen_dx = {}, set()
    rows = []
    keep = SHARED + routing + ['Diagnosis', RED_FLAGS_COL, 'Rehabilitation Phase', 'Case ID']
    n = 0
    for r in it:
        if r[0] is None:
            continue
        n += 1
        for c in counts:
            v = r[col[c]]
            counts[c][str(v) if v is not None else ''] += 1
        surgery[surgery_category(r[col['Previous Shoulder Surgery - Any (Y/N, procedure/date if known)']])] += 1
        a = r[col['Age']]
        if isinstance(a, (int, float)):
            ages.append(a)
        dx = r[col['Diagnosis']]
        rec = None
        # One real case per condition, preferring a first-presentation (Phase 1) row.
        if dx not in seen_dx or (samples[dx]['Rehabilitation Phase'] != 'Phase 1' and r[col['Rehabilitation Phase']] == 'Phase 1'):
            rec = {c: r[col[c]] for c in keep if c in col}
            samples[dx] = rec
            seen_dx.add(dx)
        if test_rows is not None:
            rows.append(r)
    test = []
    if test_rows is not None and rows:
        step = max(1, len(rows) // test_rows) if test_rows else 1
        for r in (rows[::step][:test_rows] if test_rows else rows):
            test.append({c: r[col[c]] for c in keep if c in col})
    specific = header[header.index('Cervical Spine Screen (Clear/Not Clear - referred-pain differential)') + 1:header.index('Diagnosis')]
    return {
        'file': os.path.basename(path), 'rows': n, 'columns': len(header), 'header': header, 'specific': specific,
        'counts': counts, 'surgery': surgery, 'ages': ages, 'samples': list(samples.values()), 'test': test,
    }


def jsonable(v):
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.isoformat()[:10]
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


# Multi-value cells, split the way the app splits them (engine.js splitMulti).
def multi_items(cell):
    return [x.strip() for x in re.split(r'\s*;\s*|\s*\+\s*|,\s*(?=[A-Z])', str(cell)) if x.strip()]


def separator(counter):
    """The separator a workbook uses between items of a multi-value cell, or None if it never joins items."""
    c = collections.Counter()
    for v, n in counter.items():
        v = re.sub(r'Malignancy \(specify\):[^;+]*', 'Malignancy', v)  # commas inside the specified text are not separators
        if '; ' in v:
            c['; '] += n
        if ' + ' in v:
            c[' + '] += n
        if re.search(r',\s*(?=[A-Z])', v):
            c[', '] += n
    return c.most_common(1)[0][0] if c else None


def dialect(counts):
    """How one workbook writes the shared columns, measured from its rows."""
    spec_style = collections.Counter()
    for v, n in counts['Pain_Location'].items():
        if v.startswith('Other (specify): '):
            spec_style[': '] += n
        elif v.startswith('Other (specify) - '):
            spec_style[' - '] += n
    first = lambda counter, test: next((k for k, _ in counter.most_common() if test(k)), None)
    return {
        'sep': {'Co-morbidities': separator(counts['Co-morbidities']), TREATMENT: separator(counts[TREATMENT])},
        'specifySep': {'Pain_Location': spec_style.most_common(1)[0][0] if spec_style else None},
        'notClear': first(counts[CERVICAL], lambda k: k.startswith('Not Clear')),
        'noneText': first(counts[RED_FLAGS_COL], lambda k: k.startswith('None identified')),
        'cervicalText': first(counts[RED_FLAGS_COL], lambda k: 'ervical' in k),
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--test-out', help='also write a test set of Golden Dataset rows to this path (.json, or .jsonl: one row per line)')
    ap.add_argument('--test-rows', type=int, default=400, help='rows per wrapper in the test set; 0 = every row')
    args = ap.parse_args()

    l1, ddx, grading = read_library()
    red_flags, documented = read_red_flags()
    workflow = read_workflow()

    wrappers, fields, samples, test = {}, {}, [], []
    dialects, comorb, comorb_in, free = {}, collections.Counter(), collections.defaultdict(set), {c: collections.Counter() for c in FREE_TEXT}
    shared_values = {c: collections.Counter() for c in ['Onset', 'Pain_Location']}
    jsonl = open(args.test_out, 'w', encoding='utf8') if args.test_out and args.test_out.endswith('.jsonl') else None
    for w in IDS:
        print(f'reading {w} ...', flush=True)
        d = read_dataset(w, args.test_rows if args.test_out else None)
        for c in shared_values:
            shared_values[c].update(d['counts'][c])
        dialects[w] = dialect(d['counts'])
        for v, n in d['counts']['Co-morbidities'].items():
            for item in multi_items(v):
                item = 'Malignancy (specify)' if item.startswith('Malignancy (specify)') else item
                comorb[item] += n
                comorb_in[item].add(w)
        for c in FREE_TEXT:
            free[c].update(d['counts'][c])
        flags_here = [f for f, v in red_flags.items() if w in v['wrappers']]
        rf_texts = [[k, c] for k, c in d['counts'][RED_FLAGS_COL].most_common() if k and not k.startswith('None identified')]
        wrappers[w] = {
            'id': w, 'name': grading[w]['name'], 'file': d['file'], 'rows': d['rows'], 'columns': d['columns'],
            'specificFields': d['specific'],
            'conditions': [{k: v for k, v in c.items() if k != 'rule'} for c in ddx[w]],
            'grading': {'field': grading[w]['field'], 'values': grading[w]['values']},
            'redFlags': flags_here,
            'redFlagTexts': rf_texts[:14],
            'age': {'min': min(d['ages']), 'median': statistics.median(d['ages']), 'max': max(d['ages'])} if d['ages'] else None,
            'onset': top(d['counts']['Onset']),
            'painLocation': top(d['counts']['Pain_Location']),
            'surgery': top(d['surgery']),
            'diagnoses': top(d['counts']['Diagnosis'], 20),
        }
        for c in ROUTING_FIELDS.get(w, []):
            v = l1.get((w, c), {})
            fields.setdefault(c, {})[w] = {
                'rule': v.get('rule'), 'type': v.get('type'), 'values': v.get('values'), 'notes': v.get('notes'),
                'top': top(d['counts'][c]), 'rows': d['rows'],
            }
        for s in d['samples']:
            samples.append({'wrapper': w, **{k: jsonable(v) for k, v in s.items()}})
        for s in d['test']:
            row = {'wrapper': w, **{k: jsonable(v) for k, v in s.items()}}
            if jsonl:
                jsonl.write(json.dumps(row, ensure_ascii=False) + '\n')
            else:
                test.append(row)

    shared = {}
    for c in SHARED:
        v = l1.get(('S01', c), {})
        same = all(l1.get((w, c), {}).get('values') == v.get('values') for w in IDS)
        shared[c] = {'rule': v.get('rule'), 'type': v.get('type'), 'values': v.get('values'), 'notes': v.get('notes'), 'sameInAll18': same}
    shared['Onset']['top'] = top(shared_values['Onset'], 20)
    shared['Pain_Location']['top'] = top(shared_values['Pain_Location'], 20)

    vocab = {
        'Co-morbidities': [[k, n, sorted(comorb_in[k])] for k, n in comorb.most_common()],
        'suggestions': {c: [k for k, _ in free[c].most_common(40) if k] for c in FREE_TEXT},
    }
    built = {
        'on': datetime.date.today().isoformat(),
        'rows': sum(x['rows'] for x in wrappers.values()),
        'conditions': sum(len(x['conditions']) for x in wrappers.values()),
        'sources': {
            'datasets': [x['file'] for x in wrappers.values()],
            'library': 'Shoulder_CDSS_Field_Rule_Logic_Library.xlsx (Layer1_Field_Validators, Layer3_DDX_Rules, Grading_Staging_Scale_Reference)',
            'workflow': 'Shoulder_CDSS_Workflow_and_Navigation_Guide.docx (Step 1, Step 2, section 3 table)',
            'redFlags': 'Section 10 (Red Flag Master List) of the 18 Sxx_Field_Rule_Logic_Document.docx files',
        },
    }

    parts = [
        '// GENERATED by s01/scripts/build_routing_data.py from the 18 Golden Datasets, their Field Rule Logic',
        '// Documents and the logic_docs workbooks. Do not edit by hand: change the sources and run the script again.',
        f'export const BUILT = {json.dumps(built, ensure_ascii=False)};',
        f'export const WRAPPERS = {json.dumps(wrappers, ensure_ascii=False)};',
        f'export const SHARED_FIELDS = {json.dumps(shared, ensure_ascii=False)};',
        f'export const ROUTING_FIELDS = {json.dumps(fields, ensure_ascii=False)};',
        f'export const RED_FLAGS_DOC = {json.dumps(red_flags, ensure_ascii=False)};',
        f'export const RED_FLAG_FIELD_NOTES = {json.dumps(documented, ensure_ascii=False)};',
        f'export const WORKFLOW = {json.dumps(workflow, ensure_ascii=False)};',
        f'export const DIALECT = {json.dumps(dialects, ensure_ascii=False)};',
        f'export const VOCAB = {json.dumps(vocab, ensure_ascii=False)};',
        f'export const SAMPLES = {json.dumps(samples, ensure_ascii=False)};',
        '',
    ]
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w', encoding='utf8', newline='\n') as fh:
        fh.write('\n'.join(parts))
    print(f'wrote {OUT} ({os.path.getsize(OUT) // 1024} KB): {built["rows"]} rows, {built["conditions"]} conditions, '
          f'{len(red_flags)} red flags, {len(samples)} sample cases')
    if jsonl:
        jsonl.close()
        print(f'wrote {args.test_out} (one test row per line)')
    elif args.test_out:
        with open(args.test_out, 'w', encoding='utf8') as fh:
            json.dump(test, fh, ensure_ascii=False)
        print(f'wrote {args.test_out} ({len(test)} test rows)')


if __name__ == '__main__':
    main()
